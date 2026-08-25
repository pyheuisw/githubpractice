from openpyxl import Workbook

wb = Workbook() # 새 워크북 생성
ws = wb.create_sheet()
# ws = wb.active # 현재 활성화된 sheet 가져옴
ws.sheet_properties.tabColor = "ff66ff"
ws.title = "test"

ws2 = wb.create_sheet("test2", 2)

# Sheet 복수
new_ws["A1"] = "test"
wb.save("test.xlsx")
wb.close()
